"""
Chuyển file Excel thư viện lỗi sai (tab Loi_sai) -> data/misconceptions/ham_so.csv, oxyz.csv
    pip install openpyxl
    python scripts/xlsx_to_csv.py "TutorSim_Loi_sai_56_cau.xlsx"
Cột dieu_kien_hieu_ra trong Excel: các điều kiện cách nhau bằng dấu ; (mỗi điều kiện bắt đầu bằng "Gia sư").
Giữ nguyên cột geogebra nếu CSV cũ đã có (vì Excel không có cột này).
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

OUT = Path(__file__).resolve().parents[1] / "data" / "misconceptions"
COLS = ["id", "chuyen_de", "dang_bai", "ten_loi", "bai_toan", "loi_giai_sai", "nguyen_nhan_goc", "dieu_kien_hieu_ra",
        "phan_ung_goi_y", "dap_an_dung", "geogebra", "nguon", "trang_thai", "tham_khao"]


def pipes(text, pattern):
    return " | ".join(p.strip() for p in re.split(pattern, str(text or "")) if p.strip())


def main(path):
    ws = load_workbook(path, read_only=True, data_only=True)["Loi_sai"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h else "" for h in rows[0]]
    old = {}
    for f in OUT.glob("*.csv"):
        old.update({r["id"]: r for r in csv.DictReader(open(f, encoding="utf-8"))})
    by_cd = {}
    for raw in rows[1:]:
        r = {h: ("" if v is None else str(v).strip()) for h, v in zip(head, raw)}
        if not r.get("id") or not r.get("chuyen_de"):
            continue
        dk = r.get("dieu_kien_hieu_ra", "")
        r["dieu_kien_hieu_ra"] = dk if "|" in dk else pipes(dk, r";\s*(?=Gia sư)")
        r["phan_ung_goi_y"] = pipes(r.get("phan_ung_goi_y"), r";\s+(?=[A-ZĐÂÊÔƯ])")
        r["geogebra"] = r.get("geogebra") or old.get(r["id"], {}).get("geogebra", "")
        by_cd.setdefault(r["chuyen_de"], []).append({c: r.get(c, "") for c in COLS})
    for cd, items in by_cd.items():
        with open(OUT / f"{cd}.csv", "w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=COLS)
            w.writeheader()
            w.writerows(items)
        print(f"{cd}: {len(items)} lỗi -> {OUT / (cd + '.csv')}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
