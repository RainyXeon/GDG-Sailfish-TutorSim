"""
Chuyển Excel ngân hàng câu hỏi (tab Cau_hoi) -> các file đề data/exams/*.json
    pip install openpyxl
    python scripts/questions_to_exams.py "TutorSim_Ngan_hang_cau_hoi_mau.xlsx" [--tat-ca] [--so-cau 10] [--phut 15]
Mặc định chỉ lấy câu có trang_thai = "Đồng ý". --tat-ca: lấy cả câu đang chờ duyệt (để chạy thử).
Mỗi chuyên đề chia thành các đề liên tiếp, mỗi đề --so-cau câu.
"""
import argparse
import csv
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TOPIC = {"ham_so": "Hàm số", "oxyz": "Oxyz"}


def misconception_names():
    names = {}
    for f in (ROOT / "data" / "misconceptions").glob("*.csv"):
        names.update({r["id"]: r["ten_loi"] for r in csv.DictReader(open(f, encoding="utf-8"))})
    return names


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--tat-ca", action="store_true")
    ap.add_argument("--so-cau", type=int, default=10)
    ap.add_argument("--phut", type=int, default=15)
    a = ap.parse_args()
    ws = load_workbook(a.xlsx, read_only=True, data_only=True)["Cau_hoi"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() if h else "" for h in rows[0]]
    names = misconception_names()
    by_cd = {}
    for raw in rows[1:]:
        r = {h: ("" if v is None else str(v).strip()) for h, v in zip(head, raw)}
        if not r.get("ma_cau") or (not a.tat_ca and r.get("trang_thai") != "Đồng ý"):
            continue
        mis = {k: r[f"loi_{k}"] for k in "ABCD" if r.get(f"loi_{k}")}
        by_cd.setdefault(r["chuyen_de"], []).append({
            "id": r["ma_cau"], "dang_bai": r["dang_bai"], "muc_do": r["muc_do"], "content": r["noi_dung"],
            "options": {k: r[k] for k in "ABCD"}, "answer": r["dap_an"], "misconceptions": mis, "explanation": r["giai_thich"]})
    total = sum(len(v) for v in by_cd.values())
    if total == 0:
        print("Không có câu nào. Chưa có câu 'Đồng ý'? Thử thêm --tat-ca. Giữ nguyên các đề cũ.")
        return
    out = ROOT / "data" / "exams"
    for old in out.glob("*.json"):
        old.unlink()
    n = 0
    for cd, qs in by_cd.items():
        for i in range(0, len(qs), a.so_cau):
            n += 1
            chunk = qs[i:i + a.so_cau]
            exam = {"id": f"de-{n:02d}", "title": f"Đề luyện {n:02d} – {TOPIC.get(cd, cd)}", "chuyen_de": cd,
                    "time_limit_min": a.phut, "questions": chunk,
                    "misconception_names": {m: names.get(m, m) for q in chunk for m in q["misconceptions"].values()}}
            (out / f"{exam['id']}.json").write_text(json.dumps(exam, ensure_ascii=False, indent=2), encoding="utf-8")
            print(f"{exam['id']}: {exam['title']} ({len(chunk)} câu)")


if __name__ == "__main__":
    main()
